# Regenerate the dated public dossier and blank registers from their sources.
# Development-only dependencies: Python 3, reportlab, openpyxl, and npm dependencies (tsx).
from pathlib import Path
from xml.sax.saxutils import escape
import csv
import io
import re
import subprocess
import zipfile
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

repo = Path(__file__).resolve().parents[1]
public = repo / 'public/inventory'
public.mkdir(parents=True, exist_ok=True)
font_dir = Path('/usr/share/fonts/truetype/dejavu')
for name, filename in [('Dossier', 'DejaVuSans.ttf'), ('Dossier-Bold', 'DejaVuSans-Bold.ttf')]:
    pdfmetrics.registerFont(TTFont(name, str(font_dir / filename)))
pdfmetrics.registerFontFamily('Dossier', normal='Dossier', bold='Dossier-Bold', italic='Dossier', boldItalic='Dossier-Bold')
navy = colors.HexColor('#0B1B3D')
blue = colors.HexColor('#1565D8')
styles = getSampleStyleSheet()
styles.add(ParagraphStyle(name='BodyQC', fontName='Dossier', fontSize=9, leading=13, spaceAfter=7, textColor=navy, alignment=TA_LEFT))
styles.add(ParagraphStyle(name='TitleQC', parent=styles['BodyQC'], fontName='Dossier-Bold', fontSize=18, leading=24, spaceAfter=15))
styles.add(ParagraphStyle(name='H2QC', parent=styles['BodyQC'], fontName='Dossier-Bold', fontSize=12, leading=16, textColor=blue, spaceBefore=13, spaceAfter=8, keepWithNext=True))
styles.add(ParagraphStyle(name='H3QC', parent=styles['H2QC'], fontSize=10, leading=14, spaceBefore=8))
styles.add(ParagraphStyle(name='CellQC', parent=styles['BodyQC'], fontSize=8.2, leading=11, spaceAfter=0))

def inline(text):
    text = escape(text)
    text = re.sub(r'\[([^\]]+)\]\((https?://[^)]+)\)', lambda m: f'<link href="{m.group(2)}" color="#1565D8"><u>{m.group(1)}</u></link>', text)
    text = re.sub(r'\*\*([^*]+)\*\*', r'<b>\1</b>', text)
    return text.replace('`', '')

def footer(canvas, doc):
    canvas.saveState()
    canvas.setFillColor(navy)
    canvas.setFont('Dossier', 8)
    canvas.drawString(18*mm, 13*mm, 'TAKATAK Inventory • MIEL AISSA • Dossier à compléter • 20 octobre 2026')
    canvas.drawRightString(192*mm, 13*mm, str(doc.page))
    canvas.restoreState()

source = repo / 'docs/MIEL_AISSA_REQUIS_QUEBEC.md'
lines = source.read_text().splitlines()
story = []
i = 0
while i < len(lines):
    line = lines[i].strip()
    if not line:
        i += 1
        continue
    if line.startswith('|'):
        table_rows = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r'[:\- ]+', c) for c in cells):
                table_rows.append([Paragraph(inline(c), styles['CellQC']) for c in cells])
            i += 1
        table = Table(table_rows, colWidths=[62*mm, 112*mm], repeatRows=1, hAlign='LEFT')
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#EAF2FC')),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('BOX', (0, 0), (-1, -1), .4, colors.HexColor('#CBD5E1')),
            ('INNERGRID', (0, 0), (-1, -1), .3, colors.HexColor('#CBD5E1')),
            ('LEFTPADDING', (0, 0), (-1, -1), 6),
            ('RIGHTPADDING', (0, 0), (-1, -1), 6),
            ('TOPPADDING', (0, 0), (-1, -1), 5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ]))
        story += [table, Spacer(1, 8)]
        continue
    style = 'BodyQC'
    if line.startswith('# '):
        style, line = 'TitleQC', line[2:]
    elif line.startswith('## '):
        style, line = 'H2QC', line[3:]
    elif line.startswith('### '):
        style, line = 'H3QC', line[4:]
    elif line.startswith('- '):
        line = '• ' + line[2:]
    story.append(Paragraph(inline(line), styles[style]))
    i += 1
pdf = public / 'MIEL-AISSA-requis-Quebec.pdf'
SimpleDocTemplate(str(pdf), pagesize=(210*mm, 297*mm), rightMargin=18*mm, leftMargin=18*mm, topMargin=18*mm, bottomMargin=22*mm, title='MIEL AISSA — dossier de préparation à la vente au Québec', author='TAKATAK', subject='Sources officielles et décisions à documenter; aucune approbation commerciale ou légale acquise').build(story, onFirstPage=footer, onLaterPages=footer)
registers = public / 'MIEL-AISSA-registres-Quebec.zip'
templates = repo / 'docs/inventory-quebec-templates'
catalog_process = subprocess.run(
    [str(repo / 'node_modules/.bin/tsx'), '-e',
     'import {quebecRequirements} from "./src/lib/inventory/quebec-requirements"; console.log(JSON.stringify(quebecRequirements));'],
    cwd=repo, capture_output=True, text=True, check=True,
)
catalog = json.loads(catalog_process.stdout)
book = Workbook()
intro = book.active
intro.title = 'À lire'
for row in [
    ['MIEL AISSA — registres à compléter'],
    ['Modèles vierges, 20 octobre 2026. Aucun permis ni lot réel n’est renseigné.'],
    ['Lire le dossier PDF et LIRE-AVANT.md avant de les adapter à vos opérations.'],
    ['Référencer les pièces originales et relier les lots amont, finis et livrés.'],
    ['GTIN, SKU, lots et références doivent être conservés comme texte.'],
    ['La feuille Conditionnement décrit les liens intrant/lot fini; ne pas additionner plusieurs fois la même production.'],
    ['Le suivi des pièces est déclaratif et ne délivre aucune attestation légale.'],
    ['Préparer une extraction regroupée et lisible; protéger les accès et conserver les archives requises.'],
]: intro.append(row)
intro.column_dimensions['A'].width = 115
review = book.create_sheet('Dossier des pièces')
review.append(['Rubrique', 'Nature', 'Statut', 'Référence — document, date, responsable', 'Pièces attendues'])
for index, item in enumerate(catalog, 2):
    review.append([item['title'], item['kind'], 'À vérifier', None, item['evidence']])
    options = 'À vérifier,Pièce référencée' + (',Non applicable justifié' if item['allowNotApplicable'] else '')
    validation = DataValidation(type='list', formula1='"'+options+'"', allow_blank=False)
    validation.errorTitle = 'Choisir un statut de préparation'
    validation.error = 'Une rubrique doit être documentée; une exemption doit être justifiée.'
    validation.showErrorMessage = True
    review.add_data_validation(validation)
    validation.add(f'C{index}')
for column, width in [('A', 36), ('B', 35), ('C', 28), ('D', 65), ('E', 65)]: review.column_dimensions[column].width = width
sheet_names = ['Réceptions', 'Conditionnement', 'Livraisons', 'Contrôles poids', 'Retraits réclamations']
for path, name in zip(sorted(templates.glob('*.csv')), sheet_names):
    rows = list(csv.reader(io.StringIO(path.read_text()), delimiter=';'))
    sheet = book.create_sheet(name)
    sheet.append(rows[0])
    for idx, heading in enumerate(rows[0], 1):
        letter = get_column_letter(idx)
        sheet.column_dimensions[letter].width = 28 if len(heading) < 35 else 42
        if any(token in heading.lower() for token in ['lot', 'gtin', 'sku', 'reference', 'identifiant']):
            sheet.column_dimensions[letter].number_format = '@'
            for row in range(2, 202): sheet.cell(row, idx).number_format = '@'
for sheet in book:
    sheet.freeze_panes = 'A2'
    sheet.sheet_view.showGridLines = True
    sheet.row_dimensions[1].height = 44
    for cell in sheet[1]:
        cell.fill = PatternFill('solid', fgColor='0B1B3D')
        cell.font = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    if sheet != intro:
        sheet.auto_filter.ref = f'A1:{get_column_letter(sheet.max_column)}{max(sheet.max_row, 2)}'
        for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = '1:1'
workbook = public / 'MIEL-AISSA-registres-Quebec.xlsx'
book.save(workbook)
with zipfile.ZipFile(registers, 'w', zipfile.ZIP_DEFLATED) as z:
    for path in sorted(templates.iterdir()):
        if path.suffix == '.csv':
            rows = list(csv.reader(io.StringIO(path.read_text()), delimiter=';'))
            assert len(rows) == 1 and len(rows[0]) >= 14
            data = path.read_text().encode('utf-8-sig')
        else:
            data = path.read_bytes()
        z.writestr(path.name, data)
    z.write(source, source.name)
    z.write(workbook, workbook.name)
print('Generated sourced French dossier PDF and 5 blank register templates, with no real permits, batches or customer data.')
